from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

_logger = logging.getLogger(__name__)

_XLSX_SHEET_NAME_MAX = 31  # openpyxl / Excel hard limit


def read_xlsx(path: str | Path, sheet_name: str | None = None) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    try:
        if sheet is None:
            return []
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    result = []
    for row in rows[1:]:
        item = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
        if any(value not in (None, "") for value in item.values()):
            result.append(item)
    return result


def write_xlsx(path: str | Path, sheets: dict[str, list[dict[str, Any]]]) -> None:
    workbook = Workbook()
    default = workbook.active
    if default is not None and sheets:
        workbook.remove(default)
    for sheet_name, rows in sheets.items():
        if len(sheet_name) > _XLSX_SHEET_NAME_MAX:
            _logger.warning(
                "XLSX sheet name truncated from %d to %d chars: %r → %r",
                len(sheet_name), _XLSX_SHEET_NAME_MAX,
                sheet_name, sheet_name[:_XLSX_SHEET_NAME_MAX],
            )
        sheet = workbook.create_sheet(title=sheet_name[:_XLSX_SHEET_NAME_MAX])
        if not rows:
            continue
        headers = list(rows[0].keys())
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
        for column_cells in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column_cells) + 2
            sheet.column_dimensions[column_cells[0].column_letter].width = min(width, 40)
    workbook.save(path)
