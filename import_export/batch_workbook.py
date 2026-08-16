from __future__ import annotations

"""Import the user-facing SIEKACZ batch workbook and Excel clipboard rows."""

import csv
import io
import re
import shutil
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties

from app.material_catalog import catalog_family_label


PART_SHEET_ALIASES = {"formatki", "parts", "elementy", "detale"}
STOCK_SHEET_ALIASES = {"plyty", "stock", "arkusze", "material"}

PART_COLUMNS: tuple[str, ...] = (
    "material",
    "thickness",
    "height",
    "width",
    "quantity",
    "allow_rotation",
    "priority",
    "label",
    "notes",
)
STOCK_COLUMNS: tuple[str, ...] = (
    "material",
    "thickness",
    "stock_format",
    "display_height",
    "display_width",
    "quantity",
    "stack_size",
    "priority",
    "preferred_cut_axis",
    "allow_rotation",
)

_ALIASES: dict[str, set[str]] = {
    "material": {"material", "tworzywo", "rodzajmaterialu"},
    "thickness": {"grubosc", "gruboscmm", "gr"},
    "height": {"wysokosc", "wysokoscmm", "dlugosc", "dlugoscmm", "h"},
    "width": {"szerokosc", "szerokoscmm", "w"},
    "display_height": {"wysokosc", "wysokoscmm", "h"},
    "display_width": {"szerokosc", "szerokoscmm", "w"},
    "stock_format": {"format", "formatplyty", "formatplytymm", "wymiarplyty"},
    "quantity": {"ilosc", "iloscszt", "iloscplyt", "szt", "quantity", "qty"},
    "stack_size": {"sztapel", "sztapelszt", "stack", "stacksize"},
    "priority": {"priorytet", "priority", "pierwszenstwo"},
    "preferred_cut_axis": {
        "kierunekciecia",
        "priorytetbokuciecia",
        "bokciecia",
        "cutaxis",
    },
    "allow_rotation": {"obrot", "obrotdozwolony", "obracaj", "allowrotation"},
    "label": {"etykieta", "nazwa", "label"},
    "notes": {"uwagi", "notatki", "notes"},
}


@dataclass(frozen=True)
class BatchWorkbookData:
    parts: list[dict[str, object]]
    stocks: list[dict[str, object]]


def _normalize(value: object) -> str:
    source = str(value or "").strip().casefold().translate(str.maketrans({"ł": "l", "đ": "d", "ø": "o"}))
    text = unicodedata.normalize("NFKD", source)
    ascii_text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", "", ascii_text)


def _canonical_header(value: object, allowed: Iterable[str]) -> str | None:
    normalized = _normalize(value)
    for name in allowed:
        if normalized in _ALIASES.get(name, set()):
            return name
    return None


def _number(value: object, *, row: int, field: str, positive: bool = True) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        result = float(value)
    else:
        text = str(value or "").strip().replace(" ", "").replace(",", ".")
        try:
            result = float(text)
        except ValueError as exc:
            raise ValueError(f"Wiersz {row}: pole „{field}” nie jest liczbą ({value!r}).") from exc
    if positive and result <= 0:
        raise ValueError(f"Wiersz {row}: pole „{field}” musi być większe od zera.")
    return result


def _integer(value: object, *, row: int, field: str, minimum: int = 0) -> int:
    result = _number(value, row=row, field=field, positive=False)
    if not result.is_integer() or result < minimum:
        raise ValueError(f"Wiersz {row}: pole „{field}” musi być liczbą całkowitą ≥ {minimum}.")
    return int(result)


def _boolean(value: object, default: bool = True) -> bool:
    if value is None or str(value).strip() == "":
        return default
    normalized = _normalize(value)
    if normalized in {"tak", "t", "yes", "y", "true", "1", "dozwolony"}:
        return True
    if normalized in {"nie", "n", "no", "false", "0", "zablokowany"}:
        return False
    raise ValueError(f"Nieprawidłowa wartość TAK/NIE: {value!r}.")


def _priority(value: object, *, row: int) -> int:
    if value is None or str(value).strip() == "":
        return 0
    normalized = _normalize(value)
    if normalized in {"tak", "true", "yes"}:
        return 1
    if normalized in {"nie", "false", "no"}:
        return 0
    return _integer(value, row=row, field="Priorytet", minimum=0)


def _cut_axis(value: object) -> str:
    normalized = _normalize(value)
    if not normalized or normalized in {"auto", "automatyczny", "automatycznie"}:
        return "auto"
    if normalized in {"x", "wysokosc", "wzdluzwysokosci", "pionowo"}:
        return "x"
    if normalized in {"y", "szerokosc", "wzdluzszerokosci", "poziomo"}:
        return "y"
    raise ValueError(
        f"Nieznany kierunek cięcia: {value!r}. Użyj AUTO, WZDŁUŻ WYSOKOŚCI albo WZDŁUŻ SZEROKOŚCI."
    )


def _sheet_format(value: object, *, row: int) -> tuple[float, float] | None:
    text = str(value or "").strip()
    if not text:
        return None
    match = re.fullmatch(r"\s*(\d+(?:[.,]\d+)?)\s*[x×]\s*(\d+(?:[.,]\d+)?)\s*(?:mm)?\s*", text, re.I)
    if match is None:
        raise ValueError(f"Wiersz {row}: format płyty ma mieć postać 1000×2000.")
    return (
        _number(match.group(1), row=row, field="Format płyty"),
        _number(match.group(2), row=row, field="Format płyty"),
    )


def _row_is_blank(row: dict[str, object]) -> bool:
    return not any(value not in (None, "") for value in row.values())


def _parse_part_rows(rows: list[dict[str, object]], start_row: int = 2) -> list[dict[str, object]]:
    parsed: list[dict[str, object]] = []
    for offset, row in enumerate(rows):
        row_number = start_row + offset
        if _row_is_blank(row):
            continue
        try:
            parsed.append(
                {
                    "material": str(row.get("material") or "").strip(),
                    "thickness": _number(row.get("thickness"), row=row_number, field="Grubość"),
                    "height": _number(row.get("height"), row=row_number, field="Wysokość"),
                    "width": _number(row.get("width"), row=row_number, field="Szerokość"),
                    "quantity": _integer(row.get("quantity"), row=row_number, field="Ilość", minimum=1),
                    "allow_rotation": _boolean(row.get("allow_rotation"), True),
                    "priority": _priority(row.get("priority"), row=row_number),
                    "label": str(row.get("label") or "").strip(),
                    "notes": str(row.get("notes") or "").strip(),
                }
            )
        except ValueError as exc:
            raise ValueError(f"Formatki — {exc}") from exc
    return parsed


def _parse_stock_rows(rows: list[dict[str, object]], start_row: int = 2) -> list[dict[str, object]]:
    parsed: list[dict[str, object]] = []
    for offset, row in enumerate(rows):
        row_number = start_row + offset
        if _row_is_blank(row):
            continue
        try:
            quantity = _integer(row.get("quantity"), row=row_number, field="Ilość płyt", minimum=1)
            stack_size = _integer(row.get("stack_size") or 1, row=row_number, field="Sztapel", minimum=1)
            selected_format = _sheet_format(row.get("stock_format"), row=row_number)
            if stack_size > quantity:
                raise ValueError(
                    f"Wiersz {row_number}: sztapel ({stack_size}) nie może przekraczać liczby płyt ({quantity})."
                )
            parsed.append(
                {
                    "material": str(row.get("material") or "").strip(),
                    "thickness": _number(row.get("thickness"), row=row_number, field="Grubość"),
                    # The simple UI labels the stock's first model dimension as
                    # height and the second one as width; preserve that visible
                    # order when loading the sheet.
                    "width": selected_format[0] if selected_format else _number(
                        row.get("display_height"), row=row_number, field="Wysokość"
                    ),
                    "height": selected_format[1] if selected_format else _number(
                        row.get("display_width"), row=row_number, field="Szerokość"
                    ),
                    "quantity": quantity,
                    "stack_size": stack_size,
                    "priority": _priority(row.get("priority"), row=row_number),
                    "preferred_cut_axis": _cut_axis(row.get("preferred_cut_axis")),
                    "allow_rotation": _boolean(row.get("allow_rotation"), True),
                }
            )
        except ValueError as exc:
            raise ValueError(f"Płyty — {exc}") from exc
    return parsed


def _records_from_matrix(matrix: list[list[object]], columns: tuple[str, ...]) -> tuple[list[dict[str, object]], int]:
    if not matrix:
        return [], 1
    header_index = -1
    canonical: list[str | None] = []
    for index, candidate in enumerate(matrix[:10]):
        candidate_headers = [_canonical_header(value, columns) for value in candidate]
        if sum(value is not None for value in candidate_headers) >= 3:
            header_index = index
            canonical = candidate_headers
            break
    has_header = header_index >= 0
    headers = canonical if has_header else list(columns[: len(matrix[0])])
    data_rows = matrix[header_index + 1 :] if has_header else matrix
    start_row = header_index + 2 if has_header else 1
    records: list[dict[str, object]] = []
    for values in data_rows:
        record: dict[str, object] = {}
        for index, value in enumerate(values):
            if index < len(headers) and headers[index]:
                record[str(headers[index])] = value
        records.append(record)
    return records, start_row


def _sheet_matrix(sheet) -> list[list[object]]:
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def read_batch_workbook(path: str | Path) -> BatchWorkbookData:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        parts_sheet = next(
            (sheet for sheet in workbook.worksheets if _normalize(sheet.title) in PART_SHEET_ALIASES),
            None,
        )
        stocks_sheet = next(
            (sheet for sheet in workbook.worksheets if _normalize(sheet.title) in STOCK_SHEET_ALIASES),
            None,
        )
        if parts_sheet is None and stocks_sheet is None:
            raise ValueError("Brak arkusza „Formatki” lub „Płyty”.")

        part_records, part_start = _records_from_matrix(
            _sheet_matrix(parts_sheet) if parts_sheet is not None else [], PART_COLUMNS
        )
        stock_records, stock_start = _records_from_matrix(
            _sheet_matrix(stocks_sheet) if stocks_sheet is not None else [], STOCK_COLUMNS
        )
        parts = _parse_part_rows(part_records, part_start)
        stocks = _parse_stock_rows(stock_records, stock_start)
        if not parts and not stocks:
            raise ValueError("Arkusz nie zawiera żadnych wypełnionych formatek ani płyt.")
        return BatchWorkbookData(parts=parts, stocks=stocks)
    finally:
        workbook.close()


def write_catalog_synced_template(
    template_path: str | Path,
    destination_path: str | Path,
    catalog_entries: Iterable[object],
) -> Path:
    """Create a fresh user workbook backed by the currently active catalogue.

    The shipped workbook owns the layout; this function only refreshes the
    catalogue rows, named lists, validations and compatibility formulas. It is
    deliberately called when the user opens or saves a template, so an XLSX
    never keeps stale materials after a supplier catalogue update.
    """
    source = Path(template_path)
    destination = Path(destination_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if source.resolve() != destination.resolve():
        shutil.copy2(source, destination)

    workbook = load_workbook(destination)
    catalog_sheet = workbook["Katalog"] if "Katalog" in workbook.sheetnames else workbook.create_sheet("Katalog")
    catalog_sheet.sheet_view.showGridLines = False

    records: dict[tuple[str, float, float, float], tuple[str, float, str, float]] = {}
    for entry in catalog_entries:
        try:
            material = catalog_family_label(entry).strip()
            thickness = float(getattr(entry, "thickness"))
            width = float(getattr(entry, "width", 0) or 0)
            height = float(getattr(entry, "height", 0) or 0)
            price = float(getattr(entry, "net_price_m2", 0) or 0)
        except (AttributeError, TypeError, ValueError):
            continue
        if not material or thickness <= 0:
            continue
        format_text = f"{width:g} × {height:g}" if width > 0 and height > 0 else ""
        records[(material.casefold(), thickness, width, height)] = (material, thickness, format_text, price)
    ordered = sorted(records.values(), key=lambda item: (item[0].casefold(), item[1], item[2]))
    materials = sorted({item[0] for item in ordered}, key=str.casefold)
    thicknesses = sorted({item[1] for item in ordered})
    formats = sorted({item[2] for item in ordered if item[2]}, key=lambda value: tuple(float(x) for x in re.findall(r"\d+(?:\.\d+)?", value)))

    max_clear = max(catalog_sheet.max_row, len(ordered) + 2, len(materials) + 2, len(thicknesses) + 2, len(formats) + 2)
    for row in catalog_sheet.iter_rows(min_row=2, max_row=max_clear, min_col=1, max_col=9):
        for cell in row:
            cell.value = None
    catalog_sheet["A1"], catalog_sheet["B1"], catalog_sheet["C1"], catalog_sheet["D1"] = (
        "Materiał", "Grubość [mm]", "Format płyty [mm]", "Cena netto [zł/m²]"
    )
    catalog_sheet["G1"], catalog_sheet["H1"], catalog_sheet["I1"] = "Lista materiałów", "Lista grubości", "Lista formatów"
    for row_index, values in enumerate(ordered, 2):
        for column, value in enumerate(values, 1):
            catalog_sheet.cell(row_index, column, value)
    for row_index, value in enumerate(materials, 2):
        catalog_sheet.cell(row_index, 7, value)
    for row_index, value in enumerate(thicknesses, 2):
        catalog_sheet.cell(row_index, 8, value)
    for row_index, value in enumerate(formats, 2):
        catalog_sheet.cell(row_index, 9, value)

    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in catalog_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    catalog_sheet.column_dimensions["A"].width = 34
    catalog_sheet.column_dimensions["B"].width = 15
    catalog_sheet.column_dimensions["C"].width = 22
    catalog_sheet.column_dimensions["D"].width = 20
    catalog_sheet.freeze_panes = "A2"
    catalog_sheet.auto_filter.ref = f"A1:D{max(2, len(ordered) + 1)}"
    if "KatalogSiekacz" in catalog_sheet.tables:
        catalog_sheet.tables["KatalogSiekacz"].ref = f"A1:D{max(2, len(ordered) + 1)}"

    named_lists = {
        "ListaMaterialow": ("G", max(2, len(materials) + 1)),
        "ListaGrubosci": ("H", max(2, len(thicknesses) + 1)),
        "ListaFormatow": ("I", max(2, len(formats) + 1)),
    }
    for name, (column, last_row) in named_lists.items():
        if name in workbook.defined_names:
            del workbook.defined_names[name]
        workbook.defined_names.add(DefinedName(name, attr_text=f"'Katalog'!${column}$2:${column}${last_row}"))

    def add_list_validation(sheet, cells: str, formula: str) -> None:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "Wybierz wartość z aktualnego katalogu Siekacza."
        validation.errorTitle = "Wartość spoza katalogu"
        validation.prompt = "Lista jest odświeżana z bieżącej bazy produktów przy otwieraniu szablonu."
        validation.promptTitle = "Katalog Siekacza"
        validation.showErrorMessage = True
        validation.showInputMessage = True
        sheet.add_data_validation(validation)
        validation.add(cells)

    parts = workbook["Formatki"]
    stocks = workbook["Płyty"]
    parts.data_validations.dataValidation = []
    stocks.data_validations.dataValidation = []
    add_list_validation(parts, "A3:A202", "=ListaMaterialow")
    add_list_validation(parts, "B3:B202", "=ListaGrubosci")
    add_list_validation(parts, "F3:F202", '"TAK,NIE"')
    add_list_validation(stocks, "A3:A202", "=ListaMaterialow")
    add_list_validation(stocks, "B3:B202", "=ListaGrubosci")
    add_list_validation(stocks, "C3:C202", "=ListaFormatow")
    add_list_validation(stocks, "G3:G202", '"AUTO,WZDŁUŻ WYSOKOŚCI,WZDŁUŻ SZEROKOŚCI"')
    add_list_validation(stocks, "H3:H202", '"TAK,NIE"')

    last_catalog_row = max(2, len(ordered) + 1)
    for row_index in range(3, 203):
        parts.cell(row_index, 10).value = (
            f'=IF(COUNTA(A{row_index}:I{row_index})=0,"",IF(COUNTIFS(Katalog!$A$2:$A${last_catalog_row},A{row_index},'
            f'Katalog!$B$2:$B${last_catalog_row},B{row_index})>0,"OK","BRAK MATERIAŁU / GRUBOŚCI"))'
        )
        stocks.cell(row_index, 9).value = (
            f'=IF(COUNTA(A{row_index}:H{row_index})=0,"",IF(COUNTIFS(Katalog!$A$2:$A${last_catalog_row},A{row_index},'
            f'Katalog!$B$2:$B${last_catalog_row},B{row_index},Katalog!$C$2:$C${last_catalog_row},C{row_index})>0,'
            '"OK","BRAK TAKIEJ PŁYTY W KATALOGU"))'
        )

    if workbook.calculation is None:
        workbook.calculation = CalcProperties(calcMode="auto")
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(destination)
    workbook.close()
    return destination


def parse_clipboard_rows(text: str, kind: str) -> list[dict[str, object]]:
    """Parse tab-separated Excel cells or semicolon/comma delimited text."""
    source = str(text or "").strip("\r\n")
    if not source.strip():
        return []
    delimiter = "\t" if "\t" in source else (";" if ";" in source else ",")
    matrix = [list(row) for row in csv.reader(io.StringIO(source), delimiter=delimiter)]
    if kind == "parts":
        # Preserve the long-standing quick workflow: three unlabelled Excel
        # columns mean Szerokość, Wysokość, Ilość.
        if matrix and len(matrix[0]) == 3 and not any(
            _canonical_header(value, PART_COLUMNS) for value in matrix[0]
        ):
            material, thickness = "", 1.0
            legacy = [
                {
                    "material": material,
                    "thickness": thickness,
                    "width": row[0] if len(row) > 0 else None,
                    "height": row[1] if len(row) > 1 else None,
                    "quantity": row[2] if len(row) > 2 else None,
                }
                for row in matrix
            ]
            return _parse_part_rows(legacy, 1)
        records, start = _records_from_matrix(matrix, PART_COLUMNS)
        return _parse_part_rows(records, start)
    if kind == "stocks":
        records, start = _records_from_matrix(matrix, STOCK_COLUMNS)
        return _parse_stock_rows(records, start)
    raise ValueError(f"Nieznany rodzaj tabeli: {kind!r}.")
