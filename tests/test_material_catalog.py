from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook

from app.material_catalog import catalog_family_label, catalog_revision, read_material_catalog


def test_supplier_catalog_filters_non_boards() -> None:
    path = Path(__file__).resolve().parents[1] / "sample_data" / "material_catalog.xlsx"
    items = read_material_catalog(path)
    assert len(items) >= 100
    assert any(item.material == "PA6" and item.thickness == 1 for item in items)
    pa6 = next(
        item
        for item in items
        if item.material == "PA6"
        and item.thickness == 1
        and catalog_family_label(item).endswith("NATURALNA")
    )
    assert catalog_family_label(pa6) == "PA6 PŁYTA NATURALNA"
    assert all(item.gross_price_m2 > 0 for item in items)
    assert not any(
        marker in item.product_name.upper()
        for item in items
        for marker in ("RURA", "WAŁEK", "WALEK", "PRĘT", "PRET")
    )


def test_catalog_finds_columns_and_thickness_without_fixed_sheet_name() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Nowy dostawca"
    sheet.append(["Produkt", "Jednostka", "Cena netto", "Cena brutto"])
    sheet.append(["PE 500 PŁYTA GR. 12,5 MM NATUR.", "m2", 100, 123])
    sheet.append(["PE 500 WAŁEK FI 20 MM", "mb", 50, 61.5])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "cennik.xlsx"
        workbook.save(path)
        items = read_material_catalog(path)
    assert len(items) == 1
    assert items[0].material == "PE 500 NATUR"
    assert items[0].thickness == 12.5
    assert items[0].gross_price_m2 == 123


def test_catalog_family_removes_thickness_and_normalizes_black_board_name() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Produkt", "Jednostka", "Cena brutto"])
    sheet.append(["PA6 PŁYTA GR. 8 MM CZARNY", "m2", 120])
    sheet.append(["PA6 PŁYTA GR. 12 MM CZARNA", "m2", 130])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "cennik.xlsx"
        workbook.save(path)
        items = read_material_catalog(path)
    labels = {catalog_family_label(item) for item in items}
    assert labels == {"PA6 PŁYTA CZARNA"}
    assert all("GR." not in label and "CZARNY" not in label for label in labels)


def test_catalog_family_merges_antistatic_abbreviations() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Produkt", "Jednostka", "Cena brutto"])
    sheet.append(["PŁYTA PE HD1000 GR. 8 MM CZARN/ANTYSTATYK", "m2", 120])
    sheet.append(["PŁYTA PE HD1000 GR. 12 MM CZARNA ANT", "m2", 130])
    sheet.append(["PŁYTA PE HD1000 GR. 15 MM CZARNA - ANTYSTATYCZNA", "m2", 140])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "cennik.xlsx"
        workbook.save(path)
        items = read_material_catalog(path)
    assert {catalog_family_label(item) for item in items} == {
        "PŁYTA PE HD1000 CZARNA ANTYSTATYCZNA",
    }


def test_catalog_defaults_uncoloured_pa6_and_pom_h_to_natural_board() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Produkt", "Jednostka", "Cena brutto"])
    sheet.append(["PA6 PŁYTA GR. 8 MM", "m2", 120])
    sheet.append(["POM H PŁYTA GR. 20 MM", "m2", 130])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "cennik.xlsx"
        workbook.save(path)
        items = read_material_catalog(path)
    assert {catalog_family_label(item) for item in items} == {
        "PA6 PŁYTA NATURALNA",
        "POM H PŁYTA NATURALNA",
    }


def test_matrix_price_catalog_discovers_new_rows_columns_and_material_sheets() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PP"
    sheet.append(["LP", "Grubość", "Kolor", "Format"])
    sheet.append([None, None, None, "1000x2000", "1250 x 2500", "1500x3000"])
    sheet.append([1, 10, "natual", None, 44.5, None])
    sheet.append([2, 10, "czarny", 49.0, None, 56.0])
    sheet.append([3, 12, "zielony", None, None, None])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "boral.xlsx"
        workbook.save(path)
        items = read_material_catalog(path)

    assert {
        (item.width, item.height): item.gross_price_m2
        for item in items
    } == {
        (1000.0, 2000.0): 49.0,
        (1250.0, 2500.0): 44.5,
        (1500.0, 3000.0): 56.0,
    }
    assert {catalog_family_label(item) for item in items} == {
        "PP PŁYTA NATURALNA",
        "PP PŁYTA CZARNA",
    }


def test_matrix_catalog_picks_up_later_material_colour_format_and_price_changes() -> None:
    """A live Boral workbook can grow without a hard-coded schema update."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PVC"
    sheet.append(["LP", "Grubość", "Kolor", "Format"])
    sheet.append([None, None, None, "1000x2000"])
    sheet.append([1, 5, "naturalna", "41,50 zł/m²"])

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "boral_live.xlsx"
        workbook.save(path)
        first_revision = catalog_revision(path)
        first = read_material_catalog(path)

        sheet.cell(2, 5, "1500 x 3000")
        sheet.append([2, 8, "zielona", 55.25, 61.75])
        pa6 = workbook.create_sheet("PA6G")
        pa6.append(["LP", "Grubość", "Kolor", "Format"])
        pa6.append([None, None, None, "1000×1000"])
        pa6.append([1, 12, "czarna", 77])
        workbook.save(path)

        second = read_material_catalog(path)
        assert catalog_revision(path) != first_revision

    assert len(first) == 1
    assert first[0].gross_price_m2 == 41.5
    assert {(catalog_family_label(item), item.thickness, item.width, item.height, item.gross_price_m2) for item in second} == {
        ("PVC PŁYTA NATURALNA", 5.0, 1000.0, 2000.0, 41.5),
        ("PVC PŁYTA ZIELONA", 8.0, 1000.0, 2000.0, 55.25),
        ("PVC PŁYTA ZIELONA", 8.0, 1500.0, 3000.0, 61.75),
        ("PA6G PŁYTA CZARNA", 12.0, 1000.0, 1000.0, 77.0),
    }


def test_matrix_catalog_accepts_price_per_m2_header_and_variant_column() -> None:
    """Accept the completed Boral layout alongside the original Format header."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "POM C"
    sheet.append(["LP", "Grubość", "Kolor / odmiana", "Cena netto [zł/m²] — format płyty (mm)"])
    sheet.append([None, None, None, "1000×2000", "620×3000"])
    sheet.append([1, 4, "NATUR", 297, None])
    sheet.append([2, 4, "CZARNY", 301, 305])
    sheet.append([3, 6, "NATUR", None, None])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "boral_uzupelniony.xlsx"
        workbook.save(path)
        items = read_material_catalog(path)

    assert {(catalog_family_label(item), item.thickness, item.width, item.height, item.gross_price_m2) for item in items} == {
        ("POM C PŁYTA NATURALNA", 4.0, 1000.0, 2000.0, 297.0),
        ("POM C PŁYTA CZARNA", 4.0, 1000.0, 2000.0, 301.0),
        ("POM C PŁYTA CZARNA", 4.0, 620.0, 3000.0, 305.0),
    }


if __name__ == "__main__":
    test_supplier_catalog_filters_non_boards()
    test_catalog_finds_columns_and_thickness_without_fixed_sheet_name()
    test_catalog_family_removes_thickness_and_normalizes_black_board_name()
    test_catalog_family_merges_antistatic_abbreviations()
    test_catalog_defaults_uncoloured_pa6_and_pom_h_to_natural_board()
    test_matrix_price_catalog_discovers_new_rows_columns_and_material_sheets()
    test_matrix_catalog_picks_up_later_material_colour_format_and_price_changes()
    test_matrix_catalog_accepts_price_per_m2_header_and_variant_column()
    print("test_material_catalog: OK")
