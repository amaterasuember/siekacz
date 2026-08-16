from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication

from app.simple_window import (
    PART_ALLOW_ROTATION_ROLE,
    PART_MATERIAL_COLUMN,
    PART_NOTES_ROLE,
    STOCK_MATERIAL_COLUMN,
    STOCK_STACK_COLUMN,
    SimpleCutWindow,
)
from app.material_catalog import MaterialCatalogEntry
from import_export.batch_workbook import parse_clipboard_rows, read_batch_workbook, write_catalog_synced_template


def _app() -> QApplication:
    return QApplication.instance() or QApplication([])


def test_reads_full_polish_batch_workbook() -> None:
    workbook = Workbook()
    parts = workbook.active
    parts.title = "Formatki"
    parts.append(
        [
            "Materiał",
            "Grubość [mm]",
            "Wysokość [mm]",
            "Szerokość [mm]",
            "Ilość [szt.]",
            "Obrót",
            "Priorytet",
            "Etykieta",
            "Uwagi",
        ]
    )
    parts.append(["POM-C naturalny", 18, 300, 440, 12, "NIE", 3, "A", "front"])
    stocks = workbook.create_sheet("Płyty")
    stocks.append(
        [
            "Materiał",
            "Grubość [mm]",
            "Wysokość [mm]",
            "Szerokość [mm]",
            "Ilość płyt",
            "Sztapel",
            "Priorytet",
            "Kierunek cięcia",
            "Obrót płyty",
        ]
    )
    stocks.append(["POM-C naturalny", 18, 1000, 2000, 10, 5, "TAK", "WZDŁUŻ WYSOKOŚCI", "TAK"])
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "zlecenie.xlsx"
        workbook.save(path)
        data = read_batch_workbook(path)

    assert data.parts == [
        {
            "material": "POM-C naturalny",
            "thickness": 18.0,
            "height": 300.0,
            "width": 440.0,
            "quantity": 12,
            "allow_rotation": False,
            "priority": 3,
            "label": "A",
            "notes": "front",
        }
    ]
    assert data.stocks[0]["stack_size"] == 5
    assert data.stocks[0]["priority"] == 1
    assert data.stocks[0]["preferred_cut_axis"] == "x"
    assert data.stocks[0]["width"] == 1000.0
    assert data.stocks[0]["height"] == 2000.0
    print("[OK] complete Polish workbook maps parts, stock, stack and cut direction")


def test_excel_clipboard_headers_and_stack_validation() -> None:
    parts_text = (
        "Materiał\tGrubość [mm]\tWysokość [mm]\tSzerokość [mm]\tIlość [szt.]\n"
        "PP szary\t10\t120\t300\t40"
    )
    rows = parse_clipboard_rows(parts_text, "parts")
    assert rows[0]["material"] == "PP szary"
    assert rows[0]["height"] == 120.0
    assert rows[0]["width"] == 300.0

    bad_stock = (
        "Materiał\tGrubość\tWysokość\tSzerokość\tIlość płyt\tSztapel\n"
        "PP szary\t10\t1000\t2000\t2\t3"
    )
    try:
        parse_clipboard_rows(bad_stock, "stocks")
    except ValueError as exc:
        assert "sztapel" in str(exc).casefold()
    else:
        raise AssertionError("stack larger than stock quantity should fail")
    print("[OK] Excel Ctrl+V headers parse and invalid stack is rejected")


def test_simple_window_applies_batch_metadata() -> None:
    _app()
    window = SimpleCutWindow()
    try:
        parts = [
            {
                "material": "POM-C naturalny",
                "thickness": 18,
                "height": 300,
                "width": 440,
                "quantity": 12,
                "allow_rotation": False,
                "priority": 2,
                "label": "A",
                "notes": "bez obrotu",
            }
        ]
        stocks = [
            {
                "material": "POM-C naturalny",
                "thickness": 18,
                "width": 1000,
                "height": 2000,
                "quantity": 10,
                "stack_size": 5,
                "priority": 1,
                "preferred_cut_axis": "x",
                "allow_rotation": True,
            }
        ]
        window._apply_batch_data(parts, stocks, replace=True)
        material_item = window.parts.item(0, PART_MATERIAL_COLUMN)
        assert material_item.data(PART_ALLOW_ROTATION_ROLE) is False
        assert material_item.data(PART_NOTES_ROLE) == "bez obrotu"
        assert window._stock_cell_text(0, STOCK_STACK_COLUMN) == "5"
        assert window._stock_preferred_cut_axis(0) == "x"
        assert window.stock_table.item(0, STOCK_MATERIAL_COLUMN).text() == "POM-C naturalny"

        project_parts = window._collect_parts()
        project_stock = window._collect_stock()
        assert project_parts[0].allow_rotation is False
        assert project_parts[0].priority == 2
        assert project_stock[0].stack_size == 5
        assert project_stock[0].priority == 1
    finally:
        window.close()
    print("[OK] UI keeps imported advanced fields through project collection")


def test_bundled_template_has_tables_formulas_and_validation() -> None:
    path = Path(__file__).resolve().parents[1] / "sample_data" / "SIEKACZ9000_szablon_zlecenia.xlsx"
    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook.sheetnames == ["Instrukcja", "Formatki", "Płyty", "Katalog", "Przykład"]
        parts = workbook["Formatki"]
        stocks = workbook["Płyty"]
        assert "FormatkiSiekacz" in parts.tables
        assert "PlytySiekacz" in stocks.tables
        assert str(parts["J3"].value).startswith("=IF(")
        assert str(stocks["I3"].value).startswith("=IF(")
        assert len(parts.data_validations.dataValidation) >= 5
        assert len(stocks.data_validations.dataValidation) >= 8
        assert workbook["Katalog"].max_row > 2
    finally:
        workbook.close()
    print("[OK] bundled workbook preserves editable tables, formulas and Excel validation")


def test_filled_bundled_template_roundtrips_through_importer() -> None:
    source = Path(__file__).resolve().parents[1] / "sample_data" / "SIEKACZ9000_szablon_zlecenia.xlsx"
    workbook = load_workbook(source)
    parts = workbook["Formatki"]
    stocks = workbook["Płyty"]
    parts.append([])  # keep table-managed rows intact; values are written to row 3 below
    for column, value in enumerate(["PP szary", 10, 120, 300, 40, "TAK", 2, "A", "test"], 1):
        parts.cell(3, column, value)
    for column, value in enumerate(["PP szary", 10, "1000 × 2000", 8, 4, 2, "AUTO", "TAK"], 1):
        stocks.cell(3, column, value)
    with tempfile.TemporaryDirectory() as directory:
        target = Path(directory) / "filled.xlsx"
        workbook.save(target)
        loaded = read_batch_workbook(target)
    workbook.close()
    assert loaded.parts[0]["quantity"] == 40
    assert loaded.stocks[0]["stack_size"] == 4
    assert loaded.stocks[0]["preferred_cut_axis"] == "auto"
    print("[OK] title band and table header in the shipped template import correctly")


def test_template_copy_refreshes_catalog_lists_and_exact_sheet_formats() -> None:
    source = Path(__file__).resolve().parents[1] / "sample_data" / "SIEKACZ9000_szablon_zlecenia.xlsx"
    catalog = [
        MaterialCatalogEntry("POM-C", 18, 123.45, 151.84, "POM-C PŁYTA CZARNA", width=1000, height=2000),
        MaterialCatalogEntry("PP", 10, 75.0, 92.25, "PP PŁYTA SZARA", width=1500, height=3000),
    ]
    with tempfile.TemporaryDirectory() as directory:
        target = Path(directory) / "fresh.xlsx"
        write_catalog_synced_template(source, target, catalog)
        workbook = load_workbook(target, data_only=False)
        try:
            rows = list(workbook["Katalog"].iter_rows(min_row=2, max_col=4, values_only=True))
            assert ("POM-C PŁYTA CZARNA", 18, "1000 × 2000", 123.45) in rows
            assert ("PP PŁYTA SZARA", 10, "1500 × 3000", 75) in rows
            assert "ListaMaterialow" in workbook.defined_names
            assert "ListaGrubosci" in workbook.defined_names
            assert "ListaFormatow" in workbook.defined_names
            assert "ListaMaterialow" in str(workbook["Formatki"].data_validations.dataValidation[0].formula1)
            assert "ListaFormatow" in str(workbook["Płyty"].data_validations.dataValidation[2].formula1)
        finally:
            workbook.close()
    print("[OK] each saved template refreshes materials, thicknesses, formats and prices from the active catalogue")


if __name__ == "__main__":
    test_reads_full_polish_batch_workbook()
    test_excel_clipboard_headers_and_stack_validation()
    test_simple_window_applies_batch_metadata()
    test_bundled_template_has_tables_formulas_and_validation()
    test_filled_bundled_template_roundtrips_through_importer()
    test_template_copy_refreshes_catalog_lists_and_exact_sheet_formats()
    print("BATCH WORKBOOK TESTS OK")
